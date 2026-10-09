import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from parse_external_altai import (  # noqa: E402
    join_shipments_to_panel,
    normalise_name,
    normalise_oktmo,
    parse_shipments,
    parse_tourism,
)


class ExternalAltaiParser(unittest.TestCase):
    def test_oktmo_normalisation_is_strict(self):
        self.assertEqual(normalise_oktmo("84-610-000-000"), "84610000000")
        self.assertEqual(normalise_oktmo("84 610 000 000"), "84610000000")
        with self.assertRaises(ValueError):
            normalise_oktmo("84210000000/84610000000")

    def _shipment_fixture(self, path: Path) -> None:
        workbook = Workbook()
        workbook.active.title = "Содержание"
        sheet = workbook.create_sheet("4 кв. 2024")
        sheet["C4"] = "Январь-декабрь 2024"
        rows = [
            ("Кош-Агачский муниципальный район", "84610000000", None),
            ("Всего", "101.АГ", 1250),
            ("Сельское хозяйство", "A", "…1)"),
            ("Обрабатывающие производства", "C", None),
        ]
        for index, row in enumerate(rows, start=6):
            for column, value in enumerate(row, start=1):
                sheet.cell(index, column, value)
        workbook.save(path)

    def test_suppressed_and_missing_shipments_are_not_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "shipments.xlsx"
            self._shipment_fixture(path)
            with self.assertRaisesRegex(ValueError, "Expected 11"):
                parse_shipments(path)

            # Add the other ten municipality totals to satisfy the source contract.
            workbook = Workbook()
            workbook.active.title = "Содержание"
            sheet = workbook.create_sheet("4 кв. 2024")
            sheet["C4"] = "Январь-декабрь 2024"
            row = 6
            for number in range(11):
                code = f"84{610 + number:03d}000000"
                sheet.cell(row, 1, f"Муниципалитет {number}")
                sheet.cell(row, 2, code)
                row += 1
                sheet.cell(row, 1, "Всего")
                sheet.cell(row, 2, "101.АГ")
                sheet.cell(row, 3, 100 + number)
                row += 1
                sheet.cell(row, 1, "Сектор A")
                sheet.cell(row, 2, "A")
                sheet.cell(row, 3, "...1)")
                row += 1
                sheet.cell(row, 1, "Сектор C")
                sheet.cell(row, 2, "C")
                row += 1
            workbook.save(path)
            frame, municipalities = parse_shipments(path)

        suppressed = frame.loc[frame["sector_code"] == "A"]
        missing = frame.loc[frame["sector_code"] == "C"]
        self.assertTrue(suppressed["shipment_thousand_rubles"].isna().all())
        self.assertTrue((suppressed["value_status"] == "suppressed").all())
        self.assertTrue(missing["shipment_thousand_rubles"].isna().all())
        self.assertTrue((missing["value_status"] == "missing").all())
        self.assertEqual(municipalities["oktmo"].nunique(), 11)

    def test_panel_join_requires_identical_unique_oktmo_set(self):
        codes = [f"84{610 + number:03d}000000" for number in range(11)]
        municipalities = pd.DataFrame(
            {
                "oktmo": codes,
                "shipment_municipality_name": [f"МО {n}" for n in range(11)],
            }
        )
        shipments = pd.DataFrame(
            {"oktmo": codes, "sector_code": ["TOTAL"] * 11}
        )
        panel = pd.DataFrame(
            {
                "territory_id": [str(n) for n in range(11)],
                "oktmo": codes,
                "municipal_district_name": [f"МО {n}" for n in range(11)],
            }
        )
        joined, bridge = join_shipments_to_panel(shipments, municipalities, panel)
        self.assertEqual(len(joined), 11)
        self.assertEqual(len(bridge), 11)
        bad_panel = panel.iloc[:-1].copy()
        with self.assertRaisesRegex(ValueError, "OKTMO sets differ"):
            join_shipments_to_panel(shipments, municipalities, bad_panel)

    def test_tourism_bridge_uses_exact_conservative_normalisation(self):
        self.assertEqual(normalise_name("  город\u00a0Горно-Алтайск "), "город горно-алтайск")
        self.assertNotEqual(
            normalise_name("Кош-Агачский муниципальный район"),
            normalise_name("Кош-Агачский район"),
        )
        bridge = pd.DataFrame(
            {
                "territory_id": [str(n) for n in range(11)],
                "oktmo": [f"84{610 + n:03d}000000" for n in range(11)],
                "shipment_municipality_name": [f"Муниципалитет {n}" for n in range(11)],
            }
        )
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "tourism.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "Основные показатели деятельности КСР за 2024 год"
            for index in range(11):
                row = 6 + index
                sheet.cell(row, 1, f"  МУНИЦИПАЛИТЕТ {index}  ")
                for column in range(2, 7):
                    sheet.cell(row, column, index + column)
            workbook.save(path)
            tourism, matched = parse_tourism(path, bridge)
            self.assertEqual(len(tourism), 55)
            self.assertEqual(matched["oktmo"].nunique(), 11)

            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "Основные показатели деятельности КСР за 2024 год"
            for index in range(11):
                row = 6 + index
                suffix = " лишнее" if index == 0 else ""
                sheet.cell(row, 1, f"Муниципалитет {index}{suffix}")
                for column in range(2, 7):
                    sheet.cell(row, column, index + column)
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "Unmatched non-aggregate"):
                parse_tourism(path, bridge)


if __name__ == "__main__":
    unittest.main()
